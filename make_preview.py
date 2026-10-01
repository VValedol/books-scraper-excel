"""Build preview images for the portfolio: spreadsheet views of books.xlsx and a code snippet.

Formulas are evaluated here in Python only for the preview; the workbook itself keeps live formulas.
"""

import html
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from pygments import highlight
from pygments.formatters import HtmlFormatter
from pygments.lexers import PythonLexer

HERE = Path(__file__).parent
wb = load_workbook(HERE / "books.xlsx")
rows = list(wb["Books"].iter_rows(min_row=2, values_only=True))
prices = [r[2] for r in rows]
lo, hi = min(prices), max(prices)


def price_bg(p: float) -> str:
    t = (p - lo) / (hi - lo)
    start, end = (0xE8, 0xF5, 0xE9), (0xF8, 0xB4, 0xB4)
    return "#%02x%02x%02x" % tuple(round(s + (e - s) * t) for s, e in zip(start, end))


def cell(v, style=""):
    return f'<td style="{style}">{html.escape(str(v))}</td>'


books_rows = []
for i, (cat, title, price, rating, stock, _link) in enumerate(rows[:22], start=2):
    bar = f'<div class="bar" style="width:{rating * 20}%"></div><span>{rating}</span>'
    books_rows.append(
        f'<tr><th class="rn">{i}</th>{cell(cat)}{cell(title)}'
        f'<td class="num" style="background:{price_bg(price)}">{price:.2f}</td>'
        f'<td class="rating">{bar}</td>{cell(stock, "text-align:center")}'
        f'<td><a>open</a></td></tr>'
    )

stats = defaultdict(list)
for cat, _t, price, rating, *_ in rows:
    stats[cat].append((price, rating))
summary = sorted(stats.items(), key=lambda kv: -len(kv[1]))[:16]
max_count = len(summary[0][1])
summary_rows = []
for i, (cat, items) in enumerate(summary, start=2):
    ps = [p for p, _ in items]
    rs = [r for _, r in items]
    summary_rows.append(
        f'<tr><th class="rn">{i}</th>{cell(cat)}<td class="num">{len(items)}</td>'
        f'<td class="num">{sum(ps) / len(ps):.2f}</td><td class="num">{min(ps):.2f}</td>'
        f'<td class="num">{max(ps):.2f}</td><td class="num">{sum(rs) / len(rs):.2f}</td></tr>'
    )
chart = "".join(
    f'<div class="crow"><span class="cl">{html.escape(cat)}</span>'
    f'<span class="cb" style="width:{len(items) / max_count * 100:.0f}%"></span>'
    f'<span class="cv">{len(items)}</span></div>'
    for cat, items in summary
)

code = (HERE / "scraper.py").read_text(encoding="utf-8")
start = code.index("def scrape_category")
snippet = code[start : code.index("def style_header")].rstrip()
formatter = HtmlFormatter(style="github-dark", noclasses=True, nobackground=True)
code_html = highlight(snippet, PythonLexer(), formatter)

page = f"""<!doctype html><html><head><meta charset="utf-8"><style>
body {{ font-family: Calibri, "Segoe UI", Arial, sans-serif; margin:0; padding:20px; background:#fff; font-size:13px; }}
.sheet {{ border:1px solid #c8c8c8; display:inline-block; }}
.bar-top {{ background:#217346; color:#fff; padding:6px 12px; font-family:"Segoe UI"; font-size:13px; }}
table {{ border-collapse:collapse; }}
td, th {{ border:1px solid #e1e1e1; padding:3px 8px; white-space:nowrap; height:20px; }}
thead th {{ background:#f3f3f3; color:#555; font-weight:normal; font-size:12px; }}
.hdr td {{ background:#2F6FDE; color:#fff; font-weight:bold; }}
th.rn {{ background:#f3f3f3; color:#666; font-weight:normal; width:28px; text-align:center; }}
.num {{ text-align:right; }}
.rating {{ position:relative; width:70px; }}
.rating .bar {{ position:absolute; left:0; top:3px; bottom:3px; background:linear-gradient(90deg,#F5B041,#fbe3bb); }}
.rating span {{ position:relative; padding-left:4px; }}
a {{ color:#0563C1; text-decoration:underline; }}
.tabs {{ background:#f3f3f3; border-top:1px solid #c8c8c8; padding:4px 10px; font-size:12px; }}
.tabs b {{ background:#fff; padding:3px 12px; border:1px solid #c8c8c8; border-top:none; color:#217346; }}
.tabs span {{ padding:3px 12px; color:#555; }}
.wrap {{ display:flex; gap:24px; align-items:flex-start; }}
.chart {{ border:1px solid #d0d0d0; padding:14px 16px; width:420px; }}
.chart h4 {{ margin:0 0 10px; font-weight:normal; text-align:center; font-size:15px; }}
.crow {{ display:flex; align-items:center; margin:4px 0; font-size:12px; }}
.cl {{ width:130px; text-align:right; padding-right:8px; color:#444; }}
.cb {{ height:14px; background:#4472C4; }}
.cv {{ padding-left:6px; color:#444; }}
.code {{ background:#0d1117; color:#e6edf3; border-radius:10px; padding:18px 22px; font-family:Consolas, monospace; font-size:13px; line-height:1.5; display:inline-block; }}
.code .title {{ color:#8b949e; font-size:12px; margin-bottom:8px; }}
pre {{ margin:0; }}
</style></head><body>
<div id="books" class="sheet">
  <div class="bar-top">books.xlsx — Excel</div>
  <table>
    <thead><tr><th></th><th>A</th><th>B</th><th>C</th><th>D</th><th>E</th><th>F</th></tr></thead>
    <tr class="hdr"><th class="rn">1</th><td>Category</td><td>Title</td><td>Price, GBP</td><td>Rating</td><td>In stock</td><td>Link</td></tr>
    {''.join(books_rows)}
  </table>
  <div class="tabs"><b>Books</b><span>Summary</span> &nbsp;&nbsp; 1000 rows · filters · conditional formatting</div>
</div>
<br><br>
<div id="summary" class="wrap">
  <div class="sheet">
    <div class="bar-top">books.xlsx — Excel</div>
    <table>
      <thead><tr><th></th><th>A</th><th>B</th><th>C</th><th>D</th><th>E</th><th>F</th></tr></thead>
      <tr class="hdr"><th class="rn">1</th><td>Category</td><td>Books</td><td>Avg price</td><td>Min price</td><td>Max price</td><td>Avg rating</td></tr>
      {''.join(summary_rows)}
    </table>
    <div class="tabs"><span>Books</span><b>Summary</b> &nbsp;&nbsp; formulas: COUNTIF · AVERAGEIF · MINIFS · MAXIFS</div>
  </div>
  <div class="chart"><h4>Books per category (top 16 of {len(stats)})</h4>{chart}</div>
</div>
<br><br>
<div id="code" class="code"><div class="title">scraper.py — збір категорії з пагінацією</div>{code_html}</div>
</body></html>"""

(HERE / "preview.html").write_text(page, encoding="utf-8")
print("preview.html written;", len(rows), "rows,", len(stats), "categories")
