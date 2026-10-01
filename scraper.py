"""Collect the book catalogue from books.toscrape.com into a formatted Excel report.

books.toscrape.com is a public sandbox built for scraping practice.

Usage:
    python scraper.py                     # full catalogue -> books.xlsx
    python scraper.py --out report.xlsx --delay 0.5 --max-categories 5
"""

from __future__ import annotations

import argparse
import logging
import time
from dataclasses import dataclass
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://books.toscrape.com/"
RATINGS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}
HEADER_FILL = PatternFill("solid", fgColor="2F6FDE")
HEADER_FONT = Font(bold=True, color="FFFFFF")

log = logging.getLogger("scraper")


@dataclass
class Book:
    category: str
    title: str
    price: float
    rating: int
    in_stock: bool
    url: str


class Client:
    """HTTP session with a polite delay and simple retries."""

    def __init__(self, delay: float, retries: int = 3) -> None:
        self.delay = delay
        self.retries = retries
        self.session = requests.Session()
        self.session.headers["User-Agent"] = "portfolio-demo-scraper/1.0"

    def soup(self, url: str) -> BeautifulSoup:
        for attempt in range(1, self.retries + 1):
            try:
                response = self.session.get(url, timeout=20)
                response.raise_for_status()
                response.encoding = "utf-8"
                time.sleep(self.delay)
                return BeautifulSoup(response.text, "html.parser")
            except requests.RequestException as exc:
                log.warning("Attempt %d/%d failed for %s: %s", attempt, self.retries, url, exc)
                time.sleep(attempt * 2)
        raise RuntimeError(f"Could not load {url}")


def list_categories(client: Client) -> list[tuple[str, str]]:
    page = client.soup(BASE_URL)
    links = page.select("div.side_categories ul li ul li a")
    return [(a.get_text(strip=True), urljoin(BASE_URL, a["href"])) for a in links]


def parse_price(text: str) -> float:
    return float("".join(ch for ch in text if ch.isdigit() or ch == "."))


def scrape_category(client: Client, name: str, url: str) -> list[Book]:
    books: list[Book] = []
    while url:
        page = client.soup(url)
        for card in page.select("article.product_pod"):
            link = card.select_one("h3 a")
            rating_class = card.select_one("p.star-rating")["class"]
            books.append(
                Book(
                    category=name,
                    title=link["title"],
                    price=parse_price(card.select_one("p.price_color").get_text()),
                    rating=next((RATINGS[c] for c in rating_class if c in RATINGS), 0),
                    in_stock="In stock" in card.select_one("p.availability").get_text(),
                    url=urljoin(url, link["href"]),
                )
            )
        next_link = page.select_one("li.next a")
        url = urljoin(url, next_link["href"]) if next_link else ""
    return books


def style_header(ws, columns: list[tuple[str, int]]) -> None:
    for idx, (title, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def build_workbook(books: list[Book], path: str) -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "Books"
    style_header(ws, [("Category", 22), ("Title", 60), ("Price, GBP", 12), ("Rating", 9), ("In stock", 10), ("Link", 14)])
    for book in books:
        ws.append([book.category, book.title, book.price, book.rating, "yes" if book.in_stock else "no", "open"])
        link_cell = ws.cell(row=ws.max_row, column=6)
        link_cell.hyperlink = book.url
        link_cell.style = "Hyperlink"
    last = ws.max_row
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = "0.00"
    ws.auto_filter.ref = f"A1:F{last}"
    ws.conditional_formatting.add(
        f"C2:C{last}",
        ColorScaleRule(start_type="min", start_color="E8F5E9", end_type="max", end_color="F8B4B4"),
    )
    ws.conditional_formatting.add(
        f"D2:D{last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=5, color="F5B041")
    )

    summary = wb.create_sheet("Summary")
    style_header(summary, [("Category", 22), ("Books", 9), ("Avg price", 11), ("Min price", 11), ("Max price", 11), ("Avg rating", 11)])
    categories = sorted({b.category for b in books})
    rng = f"Books!$A$2:$A${last}"
    for row, name in enumerate(categories, start=2):
        summary.cell(row=row, column=1, value=name)
        summary.cell(row=row, column=2, value=f"=COUNTIF({rng},A{row})")
        summary.cell(row=row, column=3, value=f"=ROUND(AVERAGEIF({rng},A{row},Books!$C$2:$C${last}),2)")
        summary.cell(row=row, column=4, value=f"=MINIFS(Books!$C$2:$C${last},{rng},A{row})")
        summary.cell(row=row, column=5, value=f"=MAXIFS(Books!$C$2:$C${last},{rng},A{row})")
        summary.cell(row=row, column=6, value=f"=ROUND(AVERAGEIF({rng},A{row},Books!$D$2:$D${last}),2)")
    end = len(categories) + 1
    total = end + 1
    summary.cell(row=total, column=1, value="Total").font = Font(bold=True)
    summary.cell(row=total, column=2, value=f"=SUM(B2:B{end})").font = Font(bold=True)
    summary.cell(row=total, column=3, value=f"=ROUND(AVERAGE(Books!C2:C{last}),2)").font = Font(bold=True)
    summary.auto_filter.ref = f"A1:F{end}"

    chart = BarChart()
    chart.type = "bar"
    chart.title = "Books per category"
    chart.y_axis.title = "Books"
    chart.height = 22
    chart.width = 18
    chart.legend = None
    chart.add_data(Reference(summary, min_col=2, min_row=1, max_row=end), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=1, min_row=2, max_row=end))
    summary.add_chart(chart, "H2")

    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out", default="books.xlsx", help="output .xlsx path")
    parser.add_argument("--delay", type=float, default=0.3, help="pause between requests, seconds")
    parser.add_argument("--max-categories", type=int, default=0, help="limit categories (0 = all)")
    args = parser.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    client = Client(delay=args.delay)

    categories = list_categories(client)
    if args.max_categories:
        categories = categories[: args.max_categories]
    log.info("Found %d categories", len(categories))

    books: list[Book] = []
    for name, url in categories:
        found = scrape_category(client, name, url)
        log.info("%-22s %4d books", name, len(found))
        books.extend(found)

    build_workbook(books, args.out)
    log.info("Saved %d books to %s", len(books), args.out)


if __name__ == "__main__":
    main()
